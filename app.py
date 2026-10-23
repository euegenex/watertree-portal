import os
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

EXCEL_FILE = 'feedback.xlsx'

def init_excel():
    """Initializes the Excel workbook with Watertree branding if it does not exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Equipment Queries"

        # Headers matching the 8 form sections
        headers = [
            "Log ID",
            "Timestamp",
            "Facility Name",
            "Equipment Category",
            "Brand",
            "Machine / Location",
            "Query Nature",
            "Urgency Level",
            "Issue Description",
            "Member Name",
            "Member Email",
            "Phone / WhatsApp",
            "Status"
        ]
        ws.append(headers)

        # Style Header Row (Watertree Deep Teal #0A3C42)
        header_fill = PatternFill(start_color="0A3C42", end_color="0A3C42", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color="2F9FA8"),
            right=Side(style='thin', color="2F9FA8"),
            top=Side(style='thin', color="0A3C42"),
            bottom=Side(style='medium', color="2F9FA8")
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        ws.row_dimensions[1].height = 28

        # Column widths for clean readability
        col_widths = {
            1: 12,   # Log ID
            2: 18,   # Timestamp
            3: 24,   # Facility Name
            4: 20,   # Category
            5: 18,   # Brand
            6: 25,   # Machine
            7: 24,   # Query Nature
            8: 15,   # Urgency Level
            9: 45,   # Description
            10: 20,  # Member Name
            11: 28,  # Member Email
            12: 22,  # Phone / WhatsApp
            13: 16   # Status
        }
        for col_num, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_num)].width = width

        wb.save(EXCEL_FILE)

# Ensure sheet exists on startup
init_excel()

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/submit', methods=['POST'])
def submit():
    try:
        data = request.get_json() or {}
        
        # Ensure workbook exists and load
        init_excel()
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active

        log_id = f"WT-{ws.max_row:04d}"
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        urgency = data.get('urgencyLevel', 'Medium')

        row_data = [
            log_id,
            now,
            data.get('gymLocation', 'Not Specified'),
            data.get('equipmentCategory', 'Cardio'),
            data.get('brandSelect', 'Any / Not Specified'),
            data.get('equipmentName', 'General Area'),
            data.get('queryType', 'General Query'),
            urgency,
            data.get('issueDetails', 'No details provided'),
            data.get('userName', 'Anonymous Member'),
            data.get('userEmail', 'Not Provided'),
            data.get('userPhone', 'Not Provided'),
            "New / Open"
        ]

        ws.append(row_data)
        current_row = ws.max_row

        # Standard text font and alignment
        row_font = Font(name="Calibri", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = row_font
            if col_num in [1, 2, 8, 13]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

        # Urgency highlight styling
        urgency_cell = ws.cell(row=current_row, column=8)
        if urgency == "High":
            urgency_cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
            urgency_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        elif urgency == "Medium":
            urgency_cell.font = Font(name="Calibri", size=10, bold=True, color="9C6500")
            urgency_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        ws.row_dimensions[current_row].height = 24
        wb.save(EXCEL_FILE)

        return jsonify({"status": "success", "log_id": log_id, "timestamp": now}), 200

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/excel')
def download_excel():
    """Download the full feedback.xlsx spreadsheet directly in browser."""
    if os.path.exists(EXCEL_FILE):
        return send_file(EXCEL_FILE, as_attachment=True, download_name="watertree_equipment_queries.xlsx")
    return "No feedback logged yet.", 404

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)